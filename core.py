"""Core logic: Monthly Attendance & Resource Sheet Generator.

Inputs : master Sangraha attendance workbook + one or more previous resource sheets
Outputs: per employee -> individual Sangraha sheet + updated resource sheet,
         and one combined updated resource sheet.

Column names and tab names are never trusted: every header is normalised
(lower case, punctuation/spaces removed) and matched against a list of aliases,
with a fuzzy fall-back. Column positions are always discovered, never hard-coded.
"""
import calendar
import datetime as dt
import difflib
import io
import re
from collections import Counter
from copy import copy

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RED = "FFFF0000"


# ----------------------------------------------------------------------------
# normalisation helpers
# ----------------------------------------------------------------------------
def norm(s):
    """'Rate/Month ' -> 'ratemonth'."""
    return re.sub(r"[^a-z0-9]", "", str(s).lower()) if s is not None else ""


def norm_name(s):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", str(s).lower())).strip() if s is not None else ""


def norm_id(s):
    return re.sub(r"[^A-Z0-9]", "", str(s).upper()) if s is not None else ""


MASTER_FIELDS = {
    "month_year": ["monthyear", "month", "period", "attendancemonth", "monthofattendance"],
    "vendor_id": ["vendorid", "vendorcode", "vendorno"],
    "vendor_resource_id": ["vendorresourceid", "resourceid", "employeeid", "empid", "vendorresourcecode"],
    "resource_name": ["resourcename", "employeename", "name", "resource", "empname"],
    "company": ["principalvendorcompany", "principlevendorname", "principalvendor", "vendorcompany", "company"],
    "spoc": ["spoc"],
    "hod": ["hod"],
    "present": ["presentdays", "present", "dayspresent", "daysworked"],
    "absent": ["absentdays", "absent", "daysabsent"],
    "leave": ["leavedays", "leave", "leaves", "leavetaken", "daysonleave"],
    "compoff": ["compoffdays", "compoff", "compensatoryoff", "compensatoryoffdays"],
    "holiday": ["holidaydays", "holiday", "holidays"],
    "weekoff": ["weekoffdays", "weekoff", "weekoffs", "weeklyoff", "weeklyoffdays", "weeklyoffs"],
    "total_days": ["totaldays", "totaldaysinmonth", "daysinmonth"],
}
MASTER_REQUIRED = ["resource_name", "present", "leave", "holiday", "weekoff"]

RES_FIELDS = {
    "resource_name": ["resourcename", "employeename", "name", "resource"],
    "manager": ["managername", "manager", "spoc", "reportingmanager"],
    "vendor_id": ["vendorid", "vendorresourceid", "resourceid", "employeeid", "empid"],
    "role": ["role", "designation", "position"],
    "total_leave": ["totalleave", "totalleaves", "annualleave", "leaveentitlement", "totalleaveallowed"],
    "leave_till_date": ["leavetakentilldate", "leavetakentilnow", "leavetakentillnow", "leavetilldate",
                        "leavestakentilldate", "leavetakentodate", "leavetakentilldat", "leavetaken"],
    "balance_leave": ["balanceleave", "leavebalance", "remainingleave", "balanceleaves"],
    "rate": ["ratemonth", "rate", "ratepermonth", "monthlyrate", "rates", "ratemonthly"],
    "total_billing_days": ["totalbillingdays", "totaldays", "daysinmonth", "totalbillabledays"],
    "leave_monthly": ["leavetakenmonthly", "leavemonthly", "monthlyleave", "leavethismonth", "leavetakenmonth"],
    "holiday": ["holiday", "holidays", "holidaydays"],
    "weekoff": ["weekoff", "weekoffs", "weeklyoff", "weekoffdays"],
    "days_present": ["dayspresent", "present", "presentdays", "daysworked"],
    "billing_days": ["billingdays", "billabledays"],
    "billing_amount": ["billingamount", "amount", "amounttobebilled", "billamount", "billedamount"],
    "po_number": ["ponumber", "po", "ponum", "purchaseorder", "purchaseordernumber", "pono"],
}
RES_REQUIRED = ["resource_name", "leave_till_date", "rate", "total_billing_days", "leave_monthly",
                "holiday", "weekoff", "days_present", "billing_days", "billing_amount"]

FIELD_LABELS = {
    "resource_name": "Resource Name", "present": "Present Days", "leave": "Leave Days", "holiday": "Holiday Days",
    "weekoff": "WeekOff Days", "leave_till_date": "Leave taken till date", "rate": "Rate/Month",
    "total_billing_days": "Total Billing Days", "leave_monthly": "Leave Taken monthly", "days_present": "Days Present",
    "billing_days": "Billing Days", "billing_amount": "Billing Amount", "total_leave": "Total Leave",
    "balance_leave": "Balance Leave", "po_number": "PO Number", "vendor_id": "Vendor ID",
    "vendor_resource_id": "Vendor Resource ID", "month_year": "Month/Year", "manager": "Manager Name", "role": "Role",
}


MASTER_LABELS = {
    "month_year": "Month/Year", "vendor_id": "Vendor ID", "vendor_resource_id": "Vendor Resource ID",
    "resource_name": "Resource Name", "company": "Principal Vendor Company", "spoc": "SPOC", "hod": "HOD",
    "present": "Present Days", "absent": "Absent Days", "leave": "Leave Days", "compoff": "CompOff Days",
    "holiday": "Holiday Days", "weekoff": "WeekOff Days", "total_days": "Total Days",
}


class ProcessingError(Exception):
    pass


# ----------------------------------------------------------------------------
# generic header detection
# ----------------------------------------------------------------------------
def _match_field(key, fields, claimed):
    """Return (field, fuzzy?) for a normalised header key, or (None, False)."""
    if not key:
        return None, False
    for f, aliases in fields.items():
        if f not in claimed and key in aliases:
            return f, False
    pool = {a: f for f, al in fields.items() if f not in claimed for a in al}
    if len(key) >= 5 and pool:
        hit = difflib.get_close_matches(key, list(pool), n=1, cutoff=0.86)
        if hit:
            return pool[hit[0]], True
    return None, False


def detect_header(ws, fields, max_scan=40):
    """Find the header row: the row mapping the most known fields. -> (row, {field: col}, fuzzy_notes)."""
    best = (0, None, {}, [])
    for r in range(1, min(ws.max_row, max_scan) + 1):
        claimed, mapping, notes = set(), {}, []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            f, fuzzy = _match_field(norm(v), fields, claimed)
            if f:
                claimed.add(f)
                mapping[f] = c
                if fuzzy:
                    notes.append(f"Column '{v.strip()}' was read as '{FIELD_LABELS.get(f, f)}'")
        if len(mapping) > best[0]:
            best = (len(mapping), r, mapping, notes)
    return best[1], best[2], best[3]


# ----------------------------------------------------------------------------
# value helpers
# ----------------------------------------------------------------------------
def to_num(v, what=""):
    if v is None or (isinstance(v, str) and not v.strip()):
        return 0
    if isinstance(v, bool):
        raise ValueError(f"{what}: '{v}' is not a number")
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    try:
        f = float(str(v).replace(",", "").strip())
        return int(f) if f.is_integer() else f
    except ValueError:
        raise ValueError(f"{what}: '{v}' is not a number")


MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
MONTHS.update({m.lower(): i for i, m in enumerate(calendar.month_abbr) if m})
MONTHS["sept"] = 9


def parse_month(v):
    """-> (year, month) or None."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.year, v.month
    s = str(v).strip().lower()
    if not s:
        return None
    ym = re.search(r"\b(19|20)\d\d\b", s)
    for tok in re.findall(r"[a-z]+", s):
        if tok in MONTHS:
            year = int(ym.group(0)) if ym else None
            if year is None:
                y2 = re.search(r"\b(\d\d)\b", s)
                year = 2000 + int(y2.group(1)) if y2 else None
            return (year, MONTHS[tok]) if year else None
    m = re.match(r"^(\d{4})[-/](\d{1,2})", s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r"^(\d{1,2})[-/](\d{4})$", s)
    if m:
        return int(m.group(2)), int(m.group(1))
    return None


def month_label(y, m):
    return f"{calendar.month_name[m]} {y}"


def prev_month(y, m):
    return (y - 1, 12) if m == 1 else (y, m - 1)


def ordinal(n):
    suffix = "th" if 10 <= n % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def safe_filename(s):
    return re.sub(r"[\\/:*?\"<>|]+", "-", str(s)).strip() or "unnamed"


def fmt_num(v):
    return int(v) if isinstance(v, float) and float(v).is_integer() else v


def load_wb(data, filename):
    if str(filename).lower().endswith(".xls"):
        raise ProcessingError(f"'{filename}': old .xls format is not supported - please save it as .xlsx")
    try:
        return openpyxl.load_workbook(io.BytesIO(data))
    except Exception as e:  # noqa
        raise ProcessingError(f"'{filename}' could not be opened as an Excel (.xlsx) file ({e.__class__.__name__})")


# ----------------------------------------------------------------------------
# master sheet
# ----------------------------------------------------------------------------
def parse_master(data, filename, warnings):
    wb = load_wb(data, filename)
    cands = [ws for ws in wb.worksheets if "monthly" in norm(ws.title)]
    others = [ws for ws in wb.worksheets if ws not in cands]
    chosen = None
    for pool, is_fallback in ((cands, False), (others, True)):
        scored = []
        for ws in pool:
            row, mapping, notes = detect_header(ws, MASTER_FIELDS)
            if row and all(f in mapping for f in MASTER_REQUIRED):
                scored.append((len(mapping), ws, row, mapping, notes))
        if scored:
            scored.sort(key=lambda t: -t[0])
            _, ws, row, mapping, notes = scored[0]
            chosen = (ws, row, mapping, notes, is_fallback)
            break
    if not chosen:
        raise ProcessingError(
            f"Master file '{filename}': no 'Monthly' tab with the expected columns "
            f"({', '.join(FIELD_LABELS[f] for f in MASTER_REQUIRED)}) was found. Tabs present: "
            + ", ".join(f"'{w.title}'" for w in wb.worksheets))
    ws, hrow, cmap, notes, fallback = chosen
    if fallback:
        warnings.append(f"Master file: no tab named like 'Monthly' - used tab '{ws.title}' because it has the attendance columns")
    warnings.extend(f"Master file: {n}" for n in notes)

    records = []
    for r in range(hrow + 1, ws.max_row + 1):
        name = ws.cell(r, cmap["resource_name"]).value
        if name is None or not str(name).strip():
            continue

        def g(f):
            return ws.cell(r, cmap[f]).value if f in cmap else None
        try:
            rec = {
                "row": r, "name": str(name).strip(),
                "vendor_id": g("vendor_id"), "vendor_resource_id": g("vendor_resource_id"),
                "month_raw": g("month_year"),
                "present": to_num(g("present"), f"Master row {r} Present Days"),
                "absent": to_num(g("absent"), f"Master row {r} Absent Days"),
                "leave": to_num(g("leave"), f"Master row {r} Leave Days"),
                "compoff": to_num(g("compoff"), f"Master row {r} CompOff Days"),
                "holiday": to_num(g("holiday"), f"Master row {r} Holiday Days"),
                "weekoff": to_num(g("weekoff"), f"Master row {r} WeekOff Days"),
                "total_days": to_num(g("total_days"), f"Master row {r} Total Days") if "total_days" in cmap else None,
            }
        except ValueError as e:
            raise ProcessingError(f"Master file: {e}")
        records.append(rec)
    if not records:
        raise ProcessingError(f"Master file '{filename}': tab '{ws.title}' has no employee rows")
    months = Counter(m for m in (parse_month(r["month_raw"]) for r in records) if m)
    month = months.most_common(1)[0][0] if months else None
    if len(months) > 1:
        warnings.append("Master file contains more than one month: " + ", ".join(month_label(*m) for m in months))
    headers = {f: str(ws.cell(hrow, c).value).strip() for f, c in cmap.items()}
    return {"filename": filename, "data": data, "sheet": ws.title, "header_row": hrow, "colmap": cmap,
            "headers": headers, "records": records, "month": month}


# ----------------------------------------------------------------------------
# resource sheet
# ----------------------------------------------------------------------------
def _find_label(ws, label_key, max_row):
    for r in range(1, max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and norm(v) == label_key:
                return r, c
    return None


def parse_resource(data, filename, warnings):
    wb = load_wb(data, filename)
    best = None
    for ws in wb.worksheets:
        row, mapping, notes = detect_header(ws, RES_FIELDS)
        if row and all(f in mapping for f in RES_REQUIRED):
            if not best or len(mapping) > len(best[3]):
                best = (ws, row, notes, mapping)
    if not best:
        raise ProcessingError(
            f"Resource sheet '{filename}': could not find the expected columns "
            f"({', '.join(FIELD_LABELS[f] for f in RES_REQUIRED)}). Tabs present: "
            + ", ".join(f"'{w.title}'" for w in wb.worksheets))
    ws, hrow, notes, cmap = best
    warnings.extend(f"Resource sheet '{filename}': {n}" for n in notes)
    for f in ("total_leave", "po_number", "balance_leave"):
        if f not in cmap:
            warnings.append(f"Resource sheet '{filename}': column '{FIELD_LABELS[f]}' not found - that part is skipped")

    name_col = cmap["resource_name"]
    # total row = first row below header whose name column (or neighbouring first columns) says 'Total'
    total_row, total_label_col = None, None
    for r in range(hrow + 1, ws.max_row + 1):
        for c in range(1, min(ws.max_column, name_col + 2) + 1):
            if norm(ws.cell(r, c).value) == "total":
                total_row, total_label_col = r, c
                break
        if total_row:
            break

    emp_rows_end = (total_row - 1) if total_row else ws.max_row
    employees = []
    for r in range(hrow + 1, emp_rows_end + 1):
        nm = ws.cell(r, name_col).value
        if nm is None or not str(nm).strip():
            continue

        def g(f):
            return ws.cell(r, cmap[f]).value if f in cmap else None
        employees.append({
            "src_file": filename, "src_row": r, "name": str(nm).strip(),
            "manager": g("manager"), "vendor_id": g("vendor_id"), "role": g("role"),
            "total_leave_raw": g("total_leave"), "leave_till_date_raw": g("leave_till_date"),
            "rate_raw": g("rate"), "po": g("po_number"),
        })
    if not employees:
        raise ProcessingError(f"Resource sheet '{filename}': no employee rows found under the header")

    sheet_month = None
    pos = _find_label(ws, "attendancemonth", hrow)
    if pos:
        for c in range(pos[1] + 1, ws.max_column + 1):
            if ws.cell(pos[0], c).value is not None:
                sheet_month = parse_month(ws.cell(pos[0], c).value)
                break
    if sheet_month is None:
        warnings.append(f"Resource sheet '{filename}': 'Attendance Month' not found - cannot check it is last month's sheet")

    layout = {
        "sheet": ws.title, "header_row": hrow, "colmap": cmap, "first_row": hrow + 1,
        "headers": {f: str(ws.cell(hrow, c).value).strip() for f, c in cmap.items()},
        "total_row": total_row, "total_label_col": total_label_col,
        "min_col": min(cmap.values()), "max_col": max(cmap.values()),
    }
    return {"filename": filename, "data": data, "layout": layout, "employees": employees, "month": sheet_month}


# ----------------------------------------------------------------------------
# matching
# ----------------------------------------------------------------------------
class MasterIndex:
    def __init__(self, records):
        self.records = records
        self.by_name, self.by_rid, self.by_vid = {}, {}, {}
        for i, r in enumerate(records):
            r["idx"] = i
            self.by_name.setdefault(norm_name(r["name"]), []).append(r)
            if r["vendor_resource_id"] is not None:
                self.by_rid.setdefault(norm_id(r["vendor_resource_id"]), []).append(r)
            if r["vendor_id"] is not None:
                self.by_vid.setdefault(norm_id(r["vendor_id"]), []).append(r)

    def match(self, emp):
        """-> (record, method, note) ; record None when nothing matched."""
        eid = norm_id(emp["vendor_id"])
        hits = self.by_name.get(norm_name(emp["name"]), [])
        if len(hits) > 1 and eid:
            hits = [h for h in hits if eid in (norm_id(h["vendor_resource_id"]), norm_id(h["vendor_id"]))] or hits
        if len(hits) == 1:
            return hits[0], "Resource Name", None
        why = "name appears more than once in the master" if len(hits) > 1 else "name not found in master"
        for label, idx in (("Vendor Resource ID", self.by_rid), ("Vendor ID", self.by_vid)):
            h = idx.get(eid, []) if eid else []
            if len(h) == 1:
                return h[0], label, f"{why}; matched by {label} '{emp['vendor_id']}' to master name '{h[0]['name']}'"
        return None, None, f"{why}, and ID '{emp['vendor_id']}' matched no Vendor Resource ID / Vendor ID"


# ----------------------------------------------------------------------------
# calculation
# ----------------------------------------------------------------------------
def compute(emp, rec, target):
    y, m = target
    days = calendar.monthrange(y, m)[1]
    rate = to_num(emp["rate_raw"], "Rate/Month")
    prev_leave = to_num(emp["leave_till_date_raw"], "Leave taken till date") if not (
        isinstance(emp["leave_till_date_raw"], str) and emp["leave_till_date_raw"].startswith("=")) else None
    if prev_leave is None:
        raise ValueError("Leave taken till date is a formula - cannot read its previous value")
    total_leave = None
    if emp["total_leave_raw"] not in (None, ""):
        total_leave = to_num(emp["total_leave_raw"], "Total Leave")
    new_leave = fmt_num(prev_leave + rec["leave"])
    billing_days = fmt_num(rec["leave"] + rec["holiday"] + rec["weekoff"] + rec["present"])
    amount = rate / days * billing_days if days else 0
    return {
        "days": days, "rate": rate, "prev_leave": prev_leave, "leave_till_date": new_leave,
        "total_leave": total_leave, "balance": None if total_leave is None else fmt_num(total_leave - new_leave),
        "leave_monthly": rec["leave"], "holiday": rec["holiday"], "weekoff": rec["weekoff"], "present": rec["present"],
        "billing_days": billing_days, "amount": round(amount, 2), "amount_exact": amount,
        "amount_red": round(amount, 2) < round(rate, 2),
        "row_red": total_leave is not None and new_leave > total_leave,
    }


# ----------------------------------------------------------------------------
# writers
# ----------------------------------------------------------------------------
def _font_with(font, rgb):
    f = copy(font)
    f.color = rgb
    return f


def _plain_font(font):
    """Same font, but never red (so we start from a clean baseline)."""
    col = font.color
    if col is not None and getattr(col, "type", None) == "rgb" and str(col.rgb).upper().endswith("FF0000"):
        return _font_with(font, None)
    return copy(font)


def build_sangraha(master, rec):
    wb = openpyxl.load_workbook(io.BytesIO(master["data"]))
    for ws in list(wb.worksheets):
        if ws.title != master["sheet"]:
            wb.remove(ws)
    ws = wb[master["sheet"]]
    keep = {master["header_row"], rec["row"]}
    for r in range(ws.max_row, master["header_row"], -1):
        if r not in keep:
            ws.delete_rows(r, 1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _set_date_cell(ws, row, col, y, m):
    cell = ws.cell(row, col)
    if isinstance(cell.value, (dt.datetime, dt.date)) or cell.value is None:
        cell.value = dt.datetime(y, m, 1)
    else:
        cell.value = month_label(y, m)


def _inject_cached_values(data, sheet_index, values):
    """openpyxl writes formulas without results, so viewers that do not recalculate (Excel Protected View,
    previews, mobile apps) show empty cells. Put the calculated result next to every formula we wrote."""
    import zipfile
    src = zipfile.ZipFile(io.BytesIO(data))
    name = f"xl/worksheets/sheet{sheet_index + 1}.xml"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            raw = src.read(item.filename)
            if item.filename == name:
                xml = raw.decode("utf-8")
                for coord, val in values.items():
                    xml = re.sub(r'(<c r="%s"[^>]*>\s*<f>[^<]*</f>\s*)<v\s*/>|(<c r="%s"[^>]*>\s*<f>[^<]*</f>\s*)<v></v>' % (coord, coord),
                                 lambda m: (m.group(1) or m.group(2)) + f"<v>{val}</v>", xml, count=1)
                raw = xml.encode("utf-8")
            dst.writestr(item, raw)
    return out.getvalue()


def build_resource(template, rows, target):
    """template: parsed resource dict; rows: list of (emp, rec, calc). -> xlsx bytes."""
    L = template["layout"]
    cm = L["colmap"]
    y, m = target
    wb = openpyxl.load_workbook(io.BytesIO(template["data"]))
    ws = wb[L["sheet"]]
    col = {k: get_column_letter(v) for k, v in cm.items()}
    first, total_row = L["first_row"], L["total_row"]
    n = len(rows)
    last = first + n - 1

    # baseline styles from the first employee row of the template
    base_style = {c: copy(ws.cell(first, c)._style) for c in range(L["min_col"], L["max_col"] + 1)}
    base_font = {c: _plain_font(ws.cell(first, c).font) for c in range(L["min_col"], L["max_col"] + 1)}
    total_cells = {}
    if total_row:
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(total_row, c)
            total_cells[c] = (cell.value, copy(cell._style))

    clear_to = max((total_row - 1) if total_row else last, last)
    for r in range(first, clear_to + 1):
        for c in range(L["min_col"], L["max_col"] + 1):
            cell = ws.cell(r, c)
            cell.value = None
            cell._style = copy(base_style[c])

    # month header cells
    pos = _find_label(ws, "attendancemonth", L["header_row"])
    if pos:
        for c in range(pos[1] + 1, ws.max_column + 1):
            if ws.cell(pos[0], c).value is not None:
                _set_date_cell(ws, pos[0], c, y, m)
                break
    pos = _find_label(ws, "serviceperiod", L["header_row"])
    if pos:
        cells = [c for c in range(pos[1] + 1, ws.max_column + 1) if ws.cell(pos[0], c).value is not None][:2]
        mn = calendar.month_name[m]
        if len(cells) >= 1:
            ws.cell(pos[0], cells[0]).value = f"{ordinal(1)} {mn}"
        if len(cells) >= 2:
            ws.cell(pos[0], cells[1]).value = f"{ordinal(calendar.monthrange(y, m)[1])} {mn}"

    cached = {}

    def put(r, field, value):
        if field in cm:
            ws.cell(r, cm[field]).value = value

    for i, (emp, rec, calc) in enumerate(rows):
        r = first + i
        put(r, "resource_name", emp["name"])
        put(r, "manager", emp["manager"])
        put(r, "vendor_id", emp["vendor_id"])
        put(r, "role", emp["role"])
        put(r, "total_leave", fmt_num(calc["total_leave"]) if calc["total_leave"] is not None else emp["total_leave_raw"])
        put(r, "leave_till_date", calc["leave_till_date"])
        if "total_leave" in cm and "leave_till_date" in cm:
            put(r, "balance_leave", f"={col['total_leave']}{r}-{col['leave_till_date']}{r}")
            if calc["balance"] is not None:
                cached[f"{col['balance_leave']}{r}"] = calc["balance"]
        put(r, "rate", fmt_num(calc["rate"]))
        put(r, "total_billing_days", calc["days"])
        put(r, "leave_monthly", calc["leave_monthly"])
        put(r, "holiday", calc["holiday"])
        put(r, "weekoff", calc["weekoff"])
        put(r, "days_present", calc["present"])
        put(r, "billing_days", f"={col['leave_monthly']}{r}+{col['holiday']}{r}+{col['weekoff']}{r}+{col['days_present']}{r}")
        put(r, "billing_amount", f"={col['rate']}{r}/{col['total_billing_days']}{r}*{col['billing_days']}{r}")
        cached[f"{col['billing_days']}{r}"] = calc["billing_days"]
        cached[f"{col['billing_amount']}{r}"] = calc["amount_exact"]
        put(r, "po_number", emp["po"])
        # colours: font only, never cell fill
        for c in range(L["min_col"], L["max_col"] + 1):
            cell = ws.cell(r, c)
            red = calc["row_red"] or (c == cm.get("po_number")) or (c == cm.get("billing_amount") and calc["amount_red"])
            cell.font = _font_with(base_font[c], RED) if red else copy(base_font[c])

    # total row
    if total_row:
        # stays where it was while the employees fit above it, otherwise sits right below the last employee
        # (the old total cells were already overwritten by employee rows in that case)
        new_total = total_row if last < total_row else last + 1
        for c, (val, st) in total_cells.items():
            cell = ws.cell(new_total, c)
            cell._style = copy(st)
            if isinstance(val, str) and val.upper().startswith("=SUM("):
                letter = get_column_letter(c)
                cell.value = f"=SUM({letter}{first}:{letter}{new_total - 1})"
            else:
                cell.value = val
        if "billing_amount" in cm:
            ws.cell(new_total, cm["billing_amount"]).value = f"=SUM({col['billing_amount']}{first}:{col['billing_amount']}{new_total - 1})"
            cached[f"{col['billing_amount']}{new_total}"] = sum(c["amount_exact"] for _, _, c in rows)

    ws.title = calendar.month_name[m]
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return _inject_cached_values(buf.getvalue(), wb.worksheets.index(ws), cached)


# ----------------------------------------------------------------------------
# orchestration
# ----------------------------------------------------------------------------
def process(master_file, resource_files, target_override=None):
    """master_file: (filename, bytes); resource_files: [(filename, bytes)].

    Returns {errors, warnings, employees, files:[(zip_path, bytes)], month}.
    `files` is empty when there are blocking errors.
    """
    errors, warnings = [], []
    out = {"errors": errors, "warnings": warnings, "employees": [], "files": [], "month": None}

    try:
        master = parse_master(master_file[1], master_file[0], warnings)
    except ProcessingError as e:
        errors.append(str(e))
        return out

    target = target_override or master["month"]
    if not target:
        errors.append("Could not read the month from the master's Month/Year column - please choose the target month on the page")
        return out
    out["month"] = month_label(*target)

    resources = []
    for fname, data in resource_files:
        try:
            resources.append(parse_resource(data, fname, warnings))
        except ProcessingError as e:
            errors.append(str(e))
    if not resources:
        if not errors:
            errors.append("No resource sheet was uploaded")
        return out

    for res in resources:
        pm = res["month"]
        if pm == target:
            warnings.append(f"Resource sheet '{res['filename']}' is already for {month_label(*pm)} - this month's leaves would be added a second time")
        elif pm and pm != prev_month(*target):
            warnings.append(f"Resource sheet '{res['filename']}' is for {month_label(*pm)}, not the month before {month_label(*target)}")

    index = MasterIndex(master["records"])
    seen = {}
    matched = []          # (emp, rec, calc, method, template)
    for res in resources:
        for emp in res["employees"]:
            where = f"{emp['src_file']} row {emp['src_row']}"
            rec, method, note = index.match(emp)
            if rec is None:
                errors.append(f"Employee '{emp['name']}' ({where}) could not be matched in the master by Resource Name, "
                              f"Vendor Resource ID or Vendor ID - {note}")
                continue
            if note:
                warnings.append(f"'{emp['name']}' ({where}): {note}")
            if rec["idx"] in seen:
                warnings.append(f"Employee '{emp['name']}' ({where}) is a duplicate of {seen[rec['idx']]} - using the first one")
                continue
            seen[rec["idx"]] = where
            try:
                calc = compute(emp, rec, target)
            except ValueError as e:
                errors.append(f"Employee '{emp['name']}' ({where}): {e}")
                continue
            parts = rec["present"] + rec["absent"] + rec["leave"] + rec["compoff"] + rec["holiday"] + rec["weekoff"]
            if rec["total_days"] is not None and parts != rec["total_days"]:
                warnings.append(f"'{rec['name']}': Present+Absent+Leave+CompOff+Holiday+WeekOff = {parts} but master Total Days = {rec['total_days']}")
            if rec["total_days"] is not None and rec["total_days"] != calc["days"]:
                warnings.append(f"'{rec['name']}': master Total Days = {rec['total_days']} but {month_label(*target)} has {calc['days']} days")
            if rec["compoff"]:
                warnings.append(f"'{rec['name']}' has {rec['compoff']} CompOff day(s) - CompOff is not part of Billing Days")
            if calc["row_red"]:
                warnings.append(f"'{rec['name']}': Leave taken till date ({calc['leave_till_date']}) is more than Total Leave ({calc['total_leave']}) - row shown in red")
            matched.append((emp, rec, calc, method, res))

    if errors:
        return out
    if not matched:
        errors.append("No employee could be processed")
        return out

    mlabel = f"{calendar.month_name[target[1]]}-{target[0]}"
    master_base = re.sub(r"\.xlsx?$", "", master["filename"], flags=re.I)
    used = Counter()
    files = []
    for emp, rec, calc, method, res in matched:
        nm = safe_filename(rec["name"])
        used[nm] += 1
        suffix = f" ({used[nm]})" if used[nm] > 1 else ""
        files.append((f"Sangraha - Individual/{safe_filename(master_base)} - {nm} - {mlabel}{suffix}.xlsx",
                      build_sangraha(master, rec)))
        files.append((f"Resource Sheet - Individual/Resource Sheet - {nm} - {mlabel}{suffix}.xlsx",
                      build_resource(res, [(emp, rec, calc)], target)))
    template = resources[0]
    files.append((f"Resource Sheet - All Employees - {mlabel}.xlsx",
                  build_resource(template, [(e, r, c) for e, r, c, _, _ in matched], target)))

    out["files"] = files
    out["employees"] = [{
        "name": rec["name"], "matched_by": method, "source": f"{emp['src_file']} (row {emp['src_row']})",
        "prev_leave": calc["prev_leave"], "month_leave": calc["leave_monthly"], "leave_till_date": calc["leave_till_date"],
        "total_leave": calc["total_leave"], "balance": calc["balance"], "billing_days": calc["billing_days"],
        "days": calc["days"], "rate": calc["rate"], "amount": calc["amount"],
        "amount_red": calc["amount_red"], "row_red": calc["row_red"],
    } for emp, rec, calc, method, _ in matched]
    return out


# ----------------------------------------------------------------------------
# inspection (used by the page to preview files right after upload)
# ----------------------------------------------------------------------------
def inspect_master(data, filename):
    warnings = []
    m = parse_master(data, filename, warnings)
    cols = [{"label": MASTER_LABELS[f], "header": m["headers"].get(f), "found": f in m["colmap"],
             "required": f in MASTER_REQUIRED} for f in MASTER_FIELDS]
    return {
        "filename": filename, "sheet": m["sheet"], "warnings": warnings, "columns": cols,
        "month": month_label(*m["month"]) if m["month"] else None, "month_value": (
            f"{m['month'][0]}-{m['month'][1]:02d}" if m["month"] else None),
        "employees": [{"name": r["name"], "resource_id": str(r["vendor_resource_id"] or ""),
                       "vendor_id": str(r["vendor_id"] or ""), "present": r["present"], "absent": r["absent"],
                       "leave": r["leave"], "weekoff": r["weekoff"], "holiday": r["holiday"]} for r in m["records"]],
    }


def inspect_resource(data, filename):
    warnings = []
    r = parse_resource(data, filename, warnings)
    return {
        "filename": filename, "sheet": r["layout"]["sheet"], "warnings": warnings,
        "month": month_label(*r["month"]) if r["month"] else None,
        "employees": [{"name": e["name"], "vendor_id": str(e["vendor_id"] or ""), "role": str(e["role"] or "")}
                      for e in r["employees"]],
    }
