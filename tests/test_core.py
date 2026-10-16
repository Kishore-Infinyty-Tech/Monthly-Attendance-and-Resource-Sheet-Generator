"""Run:  python tests/test_core.py   (uses the sample files in ~/Downloads)"""
import io, os, sys, copy
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl, core

D = os.path.expanduser("~/Downloads")
master = open(f"{D}/June_Infinity.xlsx", "rb").read()
rahul = open(f"{D}/Resource Sheet-Rahul-June.xlsx", "rb").read()


def red(c):
    col = c.font.color
    return col is not None and getattr(col, "rgb", None) == "FFFF0000"


def check(label, cond):
    print(("PASS " if cond else "FAIL ") + label)
    if not cond:
        check.failed = True
check.failed = False

# 1. real files, single employee ---------------------------------------------------------------------
r = core.process(("June_Infinity.xlsx", master), [("Resource Sheet-Rahul-June.xlsx", rahul)])
print("errors", r["errors"]); print("warnings", r["warnings"])
check("no errors", not r["errors"])
check("3 files (sangraha, resource, combined)", len(r["files"]) == 3)
f = dict(r["files"])
rs = openpyxl.load_workbook(io.BytesIO([v for k, v in f.items() if k.startswith("Resource Sheet - Individual")][0]))
ws = rs.active
check("tab renamed to June", ws.title == "June")
check("G5 = 3 (prev 3 + leave 0)", ws["G5"].value == 3)
check("O5 formula", ws["O5"].value == "=K5+L5+M5+N5")
check("PO red", red(ws["Q5"]))
check("amount not red", not red(ws["P5"]))
check("P19 sum", ws["P19"].value == "=SUM(P5:P18)")
sg = openpyxl.load_workbook(io.BytesIO([v for k, v in f.items() if k.startswith("Sangraha")][0]))
check("sangraha: only Monthly tab, 2 rows", sg.sheetnames == ["Monthly_June_Infinity"] and sg.active.max_row == 2
      and sg.active["D2"].value == "Rahul Desai")

# 2. changed names + multi-employee + moved total row + July target -----------------------------------
mw = openpyxl.load_workbook(io.BytesIO(master))
mws = mw["Monthly_June_Infinity"]
renames = {"Present Days": "PRESENT  days", "Resource Name": "resource name ", "Leave Days": "Leave Days.",
           "WeekOff Days": "Week Off Days", "Holiday Days": "HOLIDAYS", "Month/Year": "Month / Year"}
for c in mws[1]:
    if c.value in renames:
        c.value = renames[c.value]
mws.title = "Monthly - Infinity"
for row in mws.iter_rows(min_row=2):
    row[0].value = "July 2026"
    row[13].value = 31
    if row[3].value == "Richa B":     # 5 leaves
        pass
mw["Daily Attendance_Infinity"].title = "Weekly"
b = io.BytesIO(); mw.save(b); master2 = b.getvalue()

rw = openpyxl.load_workbook(io.BytesIO(rahul)); w = rw.active
# move total row 19 -> 8, add rows
for c in range(1, 18):
    src = w.cell(19, c); dst = w.cell(8, c)
    dst.value = src.value; dst._style = copy.copy(src._style); src.value = None
w["P8"].value = "=SUM(P5:P7)"
w["F4"].value = "Total Leaves"; w["G4"].value = "Leave Taken Till Date"; w["I4"].value = "Rate / Month"
def add(row, name, vid, rate, leave_prev, total=24):
    for c in range(2, 18):
        w.cell(row, c)._style = copy.copy(w.cell(5, c)._style)
    w.cell(row, 2).value = name; w.cell(row, 4).value = vid; w.cell(row, 5).value = "Dev"
    w.cell(row, 6).value = total; w.cell(row, 7).value = leave_prev; w.cell(row, 9).value = rate
    w.cell(row, 17).value = f"PO/{row}"
add(6, "maitri  KAUR", "TE000007", 100000, 22)       # 22 + 4 = 26 > 24 -> row red ; absent -> amount red
add(7, "Aakash S.", "TE000002", 90000, 1)           # name mismatch -> ID fallback
w["D2"].value = __import__("datetime").datetime(2026, 6, 1)
b = io.BytesIO(); rw.save(b); multi = b.getvalue()

r = core.process(("Jul master.xlsx", master2), [("multi.xlsx", multi)])
print("errors", r["errors"]); print("warnings", *r["warnings"], sep="\n  ")
check("month July 2026", r["month"] == "July 2026")
check("3 employees", len(r["employees"]) == 3)
check("7 files", len(r["files"]) == 7)
emp = {e["name"]: e for e in r["employees"]}
check("Maitri: leave 26, row red, amount red", emp["Maitri Kaur"]["leave_till_date"] == 26 and emp["Maitri Kaur"]["row_red"] and emp["Maitri Kaur"]["amount_red"])
check("Maitri: 13 billing days of 31", emp["Maitri Kaur"]["billing_days"] == 13 and emp["Maitri Kaur"]["days"] == 31)
check("Aakash matched by Vendor Resource ID", emp["Aakash Singh"]["matched_by"] == "Vendor Resource ID")
allf = dict(r["files"])
ws = openpyxl.load_workbook(io.BytesIO(allf["Resource Sheet - All Employees - July-2026.xlsx"])).active
check("combined tab July", ws.title == "July")
check("dates", str(ws["D2"].value).startswith("2026-07-01") and ws["C3"].value == "1st July" and ws["D3"].value == "31st July")
check("Maitri row red (all cells)", all(red(ws.cell(6, c)) for c in (2, 7, 16, 17)))
check("Rahul row not red except PO", not red(ws["B5"]) and red(ws["Q5"]))
check("total row stays 8", ws["B8"].value == "Total" and ws["P8"].value == "=SUM(P5:P7)")
print("PASS" if not check.failed else "SOME FAILED")
